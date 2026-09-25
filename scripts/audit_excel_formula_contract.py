#!/usr/bin/env python3
"""Validate the critical formulas of the IPS master workbook without editing it."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class FormulaContractError(RuntimeError):
    pass


def formula_text(value: Any) -> str | None:
    if isinstance(value, str) and value.startswith("="):
        return value
    text = getattr(value, "text", None)
    return text if isinstance(text, str) and text.startswith("=") else None


def audit(source: Path) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        required = {"MATRIZ", "PORTAL ESTOQUE PR", "Pesquisa Marcas", "LISTA PR-PR", "LISTA SP-SP", "LISTA PR-SC"}
        missing_sheets = sorted(required.difference(workbook.sheetnames))
        if missing_sheets:
            raise FormulaContractError("ABAS_AUSENTES:" + ",".join(missing_sheets))

        matrix = workbook["MATRIZ"]
        last_product_row = 2
        for row_number, (product_code,) in enumerate(
            matrix.iter_rows(min_row=3, min_col=1, max_col=1, values_only=True),
            start=3,
        ):
            if product_code not in (None, "", 0):
                last_product_row = row_number
        if last_product_row < 3:
            raise FormulaContractError("MATRIZ_SEM_PRODUTOS")

        search = workbook["Pesquisa Marcas"]
        search_formula = formula_text(search["A10"].value)
        counter_formula = formula_text(search["K4"].value)
        if not search_formula:
            raise FormulaContractError("FORMULA_CRITICA_AUSENTE:Pesquisa Marcas!A10")
        search_tokens = ("FILTER", "XLOOKUP", "PORTAL ESTOQUE PR", "MATRIZ")
        missing_tokens = [token for token in search_tokens if token not in search_formula.upper()]
        if missing_tokens:
            raise FormulaContractError("FORMULA_BUSCA_INCOMPLETA:Pesquisa Marcas!A10:" + ",".join(missing_tokens))
        if not counter_formula:
            raise FormulaContractError("FORMULA_CRITICA_AUSENTE:Pesquisa Marcas!K4")

        stock = workbook["PORTAL ESTOQUE PR"]
        last_stock_row = 1
        for row_number, (product_code,) in enumerate(
            stock.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True),
            start=2,
        ):
            if product_code not in (None, "", 0):
                last_stock_row = row_number
        stock_lookup_rows = [
            int(match)
            for match in re.findall(
                r"'PORTAL ESTOQUE PR'!\$[BH]\$2:\$[BH]\$(\d+)",
                search_formula,
                flags=re.IGNORECASE,
            )
        ]
        warnings: list[str] = []
        portal_segments = [
            segment.strip()[:500]
            for segment in search_formula.split(",")
            if "PORTAL ESTOQUE PR" in segment.upper()
        ]
        lookup_last_row = min(stock_lookup_rows) if stock_lookup_rows else None
        if lookup_last_row is None:
            warnings.append("FORMULA_BUSCA_SEM_INTERVALO_ESTOQUE_PADRAO:Pesquisa Marcas!A10")
        elif lookup_last_row < last_stock_row:
            warnings.append(
                f"INTERVALO_ESTOQUE_DESATUALIZADO:Pesquisa Marcas!A10:{lookup_last_row}:{last_stock_row}"
            )

        list_summary: dict[str, Any] = {}
        for sheet_name in ("LISTA PR-PR", "LISTA SP-SP", "LISTA PR-SC"):
            sheet = workbook[sheet_name]
            gaps: list[str] = []
            formula_count = 0
            for cells in sheet.iter_rows(
                min_row=3,
                max_row=last_product_row,
                min_col=1,
                max_col=11,
            ):
                for cell in cells:
                    if formula_text(cell.value):
                        formula_count += 1
                    elif len(gaps) < 20:
                        gaps.append(cell.coordinate)
            if gaps:
                raise FormulaContractError(
                    f"FORMULAS_LISTA_AUSENTES:{sheet_name}:" + ",".join(gaps)
                )
            list_summary[sheet_name] = {
                "first_row": 3,
                "last_row": last_product_row,
                "formula_count": formula_count,
            }

        return {
            "status": "OK",
            "last_product_row": last_product_row,
            "search_formula": "Pesquisa Marcas!A10",
            "counter_formula": "Pesquisa Marcas!K4",
            "last_stock_row": last_stock_row,
            "stock_lookup_last_row": lookup_last_row,
            "portal_formula_segments": portal_segments,
            "warnings": warnings,
            "lists": list_summary,
        }
    finally:
        workbook.close()
