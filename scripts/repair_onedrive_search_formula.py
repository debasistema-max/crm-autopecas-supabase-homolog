#!/usr/bin/env python3
"""Repair only the stock range in Pesquisa Marcas!A10 on personal OneDrive.

The Microsoft access token is short-lived, kept only in memory and never stored.
The upload is conditional on the workbook eTag so a concurrent edit is not lost.
"""

from __future__ import annotations

import argparse
import json
import posixpath
import re
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook

import audit_excel_formula_contract as formula_audit
import sync_onedrive_personal as onedrive


AUTHORITY = "https://login.microsoftonline.com/consumers/oauth2/v2.0"
GRAPH_ROOT = "https://graph.microsoft.com/v1.0"
SCOPES = "Files.ReadWrite"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WORKBOOK_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def post_form(url: str, data: dict[str, str]) -> dict:
    request = urllib.request.Request(
        url,
        data=urllib.parse.urlencode(data).encode("ascii"),
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            return json.loads(response.read())
    except urllib.error.HTTPError as error:
        body = json.loads(error.read() or b"{}")
        body["_status"] = error.code
        return body


def authorize(client_id: str) -> str:
    device = post_form(
        f"{AUTHORITY}/devicecode",
        {"client_id": client_id, "scope": SCOPES},
    )
    if not device.get("device_code"):
        raise RuntimeError(f"DISPOSITIVO_NAO_AUTORIZADO:{device.get('error', 'UNKNOWN')}")
    print("Abra:", device.get("verification_uri", "https://microsoft.com/devicelogin"), flush=True)
    print("Código:", device.get("user_code", ""), flush=True)
    print(
        "Entre com debasistema@gmail.com e autorize a edição única da planilha.",
        flush=True,
    )

    deadline = time.monotonic() + int(device.get("expires_in", 900))
    interval = max(int(device.get("interval", 5)), 5)
    while time.monotonic() < deadline:
        time.sleep(interval)
        token = post_form(
            f"{AUTHORITY}/token",
            {
                "grant_type": "urn:ietf:params:oauth:grant-type:device_code",
                "client_id": client_id,
                "device_code": str(device["device_code"]),
            },
        )
        if token.get("access_token"):
            return str(token["access_token"])
        error = token.get("error")
        if error == "authorization_pending":
            continue
        if error == "slow_down":
            interval += 5
            continue
        raise RuntimeError(f"AUTORIZACAO_RECUSADA:{error or 'UNKNOWN'}")
    raise RuntimeError("AUTORIZACAO_EXPIRADA")


def worksheet_part(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    sheet = next(
        (
            node
            for node in workbook_root.findall(f".//{{{WORKBOOK_NS}}}sheet")
            if node.attrib.get("name") == sheet_name
        ),
        None,
    )
    if sheet is None:
        raise RuntimeError(f"ABA_AUSENTE:{sheet_name}")
    relationship_id = sheet.attrib.get(f"{{{OFFICE_REL_NS}}}id")
    relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationship = next(
        (
            node
            for node in relationships.findall(f".//{{{PACKAGE_REL_NS}}}Relationship")
            if node.attrib.get("Id") == relationship_id
        ),
        None,
    )
    if relationship is None:
        raise RuntimeError(f"RELACIONAMENTO_ABA_AUSENTE:{sheet_name}")
    target = str(relationship.attrib.get("Target") or "")
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join("xl", target))


def last_used_stock_row(source: Path) -> int:
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook["PORTAL ESTOQUE PR"]
        last_row = 1
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=min(max(sheet.max_column, 1), 32), values_only=True),
            start=2,
        ):
            if any(value not in (None, "") for value in values):
                last_row = row_number
        if last_row < 2:
            raise RuntimeError("PORTAL_ESTOQUE_PR_SEM_DADOS")
        return last_row
    finally:
        workbook.close()


def repair_workbook(source: Path, destination: Path) -> dict[str, object]:
    stock_last_row = last_used_stock_row(source)
    cell_pattern = re.compile(rb'<c\b(?=[^>]*\br="A10")[^>]*>.*?</c>', re.DOTALL)
    reference_pattern = re.compile(
        rb"('PORTAL ESTOQUE PR'!\$[A-Z]{1,3}\$2:\$[A-Z]{1,3}\$)(\d+)",
        re.IGNORECASE,
    )

    with zipfile.ZipFile(source, "r") as source_archive:
        search_part = worksheet_part(source_archive, "Pesquisa Marcas")
        sheet_xml = source_archive.read(search_part)
        cell_matches = list(cell_pattern.finditer(sheet_xml))
        if len(cell_matches) != 1:
            raise RuntimeError("CELULA_BUSCA_A10_NAO_LOCALIZADA")
        current_cell = cell_matches[0].group(0)
        previous_rows = [
            int(last_row)
            for _prefix, last_row in reference_pattern.findall(current_cell)
        ]
        if len(previous_rows) < 2:
            raise RuntimeError("INTERVALOS_ESTOQUE_NAO_LOCALIZADOS_EM_A10")
        replacement_row = str(stock_last_row).encode("ascii")
        updated_cell, replacement_count = reference_pattern.subn(
            lambda match: match.group(1) + replacement_row,
            current_cell,
        )
        updated_cell = re.sub(rb"<v(?:\s[^>]*)?>.*?</v>|<v\s*/>", b"", updated_cell, flags=re.DOTALL)
        if replacement_count < 2:
            raise RuntimeError("INTERVALOS_ESTOQUE_INCOMPLETOS_EM_A10")
        updated_sheet_xml = sheet_xml[: cell_matches[0].start()] + updated_cell + sheet_xml[cell_matches[0].end() :]

        with zipfile.ZipFile(destination, "w") as target_archive:
            for info in source_archive.infolist():
                payload = updated_sheet_xml if info.filename == search_part else source_archive.read(info.filename)
                target_archive.writestr(info, payload)

    with zipfile.ZipFile(destination, "r") as check_archive:
        corrupt = check_archive.testzip()
        if corrupt:
            raise RuntimeError(f"XLSX_CORROMPIDO:{corrupt}")
    audit = formula_audit.audit(destination)
    if audit.get("stock_lookup_last_row") != stock_last_row:
        raise RuntimeError("FORMULA_CORRIGIDA_NAO_COBRE_ESTOQUE_ATUAL")
    stale_warnings = [
        warning
        for warning in audit.get("warnings", [])
        if str(warning).startswith("INTERVALO_ESTOQUE_DESATUALIZADO")
    ]
    if stale_warnings:
        raise RuntimeError("FORMULA_CORRIGIDA_CONTINUA_DESATUALIZADA")
    return {
        "sheet": "Pesquisa Marcas",
        "cell": "A10",
        "previous_last_rows": previous_rows,
        "new_last_row": stock_last_row,
        "replacements": replacement_count,
    }


def upload_workbook(token: str, item: dict, source: Path) -> dict:
    item_id = urllib.parse.quote(str(item.get("id") or ""), safe="")
    if not item_id:
        raise RuntimeError("ITEM_ONEDRIVE_SEM_ID")
    request = urllib.request.Request(
        f"{GRAPH_ROOT}/me/drive/items/{item_id}/content",
        data=source.read_bytes(),
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": CONTENT_TYPE,
            "If-Match": str(item.get("eTag") or ""),
        },
        method="PUT",
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            return json.loads(response.read())
    except urllib.error.HTTPError as error:
        if error.code == 412:
            raise RuntimeError("PLANILHA_ALTERADA_DURANTE_CORRECAO") from error
        raise RuntimeError(f"UPLOAD_ONEDRIVE_HTTP_{error.code}") from error


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--client-id", required=True)
    parser.add_argument("--folder-path", required=True)
    parser.add_argument("--workbook-name", required=True)
    args = parser.parse_args()

    token = authorize(args.client_id)
    before = onedrive.locate_workbook(token, args.workbook_name, args.folder_path)
    with tempfile.TemporaryDirectory(prefix="ips-formula-repair-") as temp_dir:
        original = Path(temp_dir, "original.xlsx")
        repaired = Path(temp_dir, "repaired.xlsx")
        onedrive.download_workbook(token, before, original)
        repair = repair_workbook(original, repaired)
        uploaded = upload_workbook(token, before, repaired)
    print(
        json.dumps(
            {
                "ok": True,
                "repair": repair,
                "remote_name": uploaded.get("name"),
                "remote_size": uploaded.get("size"),
                "remote_updated_at": uploaded.get("lastModifiedDateTime"),
            },
            ensure_ascii=False,
            separators=(",", ":"),
        )
    )


if __name__ == "__main__":
    main()
