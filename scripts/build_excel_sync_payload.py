#!/usr/bin/env python3
"""Build the values-only Data Sync v1 payload from the IPS master workbook.

The script never edits the workbook and opens it with data_only=True, so formula
text is not exported. Excel (or the upstream workbook service) must calculate
and save the workbook before this adapter reads it.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import unicodedata
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def sha256_file(source: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    if not hasattr(source, "open"):
        digest.update(source.read_bytes())
        return digest.hexdigest()
    with source.open("rb") as stream:
        for chunk in iter(lambda: stream.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def key(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-z0-9%]+", " ", text.lower()).strip()


def product_code(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if math.isfinite(value) and value.is_integer() else None
    text = re.sub(r"\s+", "", str(value).replace("\u200b", "").replace("\ufeff", ""))
    if re.fullmatch(r"\d+[.,]0+", text):
        return re.sub(r"[.,]0+$", "", text)
    if re.fullmatch(r"\d+(?:[.,]\d+)?[eE]\+?\d+", text):
        try:
            number = Decimal(text.replace(",", "."))
            return str(int(number)) if number == number.to_integral_value() else None
        except InvalidOperation:
            return None
    return text or None


def number(value: Any) -> float | int | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value if math.isfinite(float(value)) else None
    text = re.sub(r"[^0-9,\.\-]", "", str(value))
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        parsed = float(text)
        return int(parsed) if parsed.is_integer() else parsed
    except ValueError:
        return None


def money(value: Any) -> float | None:
    """Normalize workbook currency values to cents before JSON serialization."""
    parsed = number(value)
    if parsed is None:
        return None
    return float(Decimal(str(parsed)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def rate(value: Any) -> float | None:
    parsed = number(value)
    if parsed is None:
        return None
    return float(parsed) / 100 if "%" in str(value) or abs(float(parsed)) > 1 else float(parsed)


def workbook_fraction_rate(value: Any) -> float | None:
    """Read rates from fiscal sheets whose numeric cells store Excel fractions.

    A formatted 156% cell is stored as 1.56, not 156. Text containing a percent
    sign still needs conversion because it is not an Excel numeric fraction.
    """
    parsed = number(value)
    if parsed is None:
        return None
    return float(parsed) / 100 if "%" in str(value) else float(parsed)


def text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def digits(value: Any) -> str | None:
    result = re.sub(r"\D", "", str(value or ""))
    return result or None


def yes_no(value: Any) -> bool:
    normalized = key(value)
    return normalized in {"sim", "s", "yes", "y", "true", "1"}


def find_header(ws, required: set[str], limit: int = 12):
    for row in ws.iter_rows(min_row=1, max_row=min(limit, ws.max_row), values_only=True):
        mapping: dict[str, int] = {}
        for index, value in enumerate(row):
            if key(value):
                mapping.setdefault(key(value), index)
        if required.issubset(mapping):
            return row, mapping
    raise ValueError(f"Cabeçalho não encontrado em {ws.title}: {sorted(required)}")


def row_dict(row: tuple[Any, ...], mapping: dict[str, int]) -> dict[str, Any]:
    return {name: row[index] if index < len(row) else None for name, index in mapping.items()}


def put(fields: dict[str, Any], name: str, value: Any):
    if value is not None and value != "":
        fields[name] = value


def product_record(store: dict[str, dict[str, Any]], code: str) -> dict[str, Any]:
    return store.setdefault(code, {"area": "PRODUCT", "product_code": code, "fields": {}})


def read_products(workbook, products: dict[str, dict[str, Any]]):
    ws = workbook["MATRIZ"]
    _, mapping = find_header(ws, {"codigo ips", "descricao", "preco s imp"})
    for values in ws.iter_rows(min_row=3, values_only=True):
        row = row_dict(values, mapping)
        code = product_code(row.get("codigo ips"))
        if not code:
            continue
        fields = product_record(products, code)["fields"]
        put(fields, "description", text(row.get("descricao")))
        put(fields, "brand", text(row.get("marca")))
        put(fields, "application", text(row.get("aplicacao")))
        put(fields, "year", text(row.get("ano")))
        put(fields, "ipi_rate", rate(row.get("ipi")))

    ws = workbook["Cadastro Item SAP"]
    _, mapping = find_header(ws, {"n do item", "descricao do item", "codigo ncm"})
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = row_dict(values, mapping)
        code = product_code(row.get("n do item"))
        if not code:
            continue
        fields = product_record(products, code)["fields"]
        if not fields.get("description"):
            put(fields, "description", text(row.get("descricao do item")))
        if not fields.get("brand"):
            put(fields, "brand", text(row.get("marca")))
        put(fields, "model", text(row.get("modelo")))
        put(fields, "year", fields.get("year") or text(row.get("ano")))
        put(fields, "ncm", text(row.get("codigo ncm")))
        cest = row.get("codigo cest")
        put(fields, "cest", text(cest))
        put(fields, "ipi_rate", fields.get("ipi_rate") if fields.get("ipi_rate") is not None else rate(row.get("ipi")))
        put(fields, "oem_01", text(row.get("oem 01")))
        put(fields, "item_group", text(row.get("grupo de itens")))
        put(fields, "sales_unit", text(row.get("um venda")))
        put(fields, "item_notes", text(row.get("obs item")))
        put(fields, "weight", number(row.get("peso")))
        put(fields, "volume", number(row.get("volume")))
        put(fields, "manufacturer", text(row.get("fabricante")))
        put(fields, "barcode", text(row.get("codigo de barras")))
        put(fields, "origin_description", text(row.get("fonte do produto")))
        put(fields, "material_group", text(row.get("grupo de materiais")))
        put(fields, "fiscal_group", text(row.get("origem e grupo fiscal")))

    if "NCM Produtos" in workbook.sheetnames:
        ws = workbook["NCM Produtos"]
        _, mapping = find_header(ws, {"codigo", "ncm", "status"})
        header_row = next(i for i, row in enumerate(ws.iter_rows(values_only=True), 1)
                          if {"codigo", "ncm", "status"}.issubset({key(v) for v in row}))
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            row = row_dict(values, mapping)
            code = product_code(row.get("codigo"))
            if not code:
                continue
            fields = product_record(products, code)["fields"]
            put(fields, "ncm", text(row.get("ncm")))
            put(fields, "cest", text(row.get("cest")))
            put(fields, "ipi_rate", rate(row.get("ipi sap")))
            put(fields, "origin_description", text(row.get("origem")))


def read_stock(workbook, records: list[dict[str, Any]]):
    for branch in ("PR", "SP"):
        sheet = f"PORTAL ESTOQUE {branch}"
        if sheet not in workbook.sheetnames or workbook[sheet].max_row < 2:
            continue
        ws = workbook[sheet]
        try:
            _, mapping = find_header(ws, {"codigo", "estoque", "disp geral"})
        except ValueError:
            continue
        header_row = 1
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            row = row_dict(values, mapping)
            code = product_code(row.get("codigo"))
            if not code:
                continue
            fields: dict[str, Any] = {}
            for source, target in (
                ("estoque", "stock_qty"), ("confirmado", "confirmed_qty"),
                ("disp venda", "sales_available_qty"), ("qtd auth pend", "authorized_pending_qty"),
                ("disp geral", "general_available_qty")
            ):
                raw = row.get(source)
                parsed = number(raw)
                put(fields, target, parsed)
                if target == "general_available_qty" and raw is not None:
                    display = str(raw).strip()
                    put(fields, "source_display_value", display)
                    if "+" in display:
                        fields["general_available_capped"] = True
            if fields:
                records.append({"area": "STOCK", "product_code": code, "branch_code": branch,
                                "fields": fields, "field_mask": list(fields)})


def read_prices(workbook, records: list[dict[str, Any]]):
    calculations: dict[tuple[str, str], dict[str, Any]] = {}
    if "Cálculo Fiscal" in workbook.sheetnames:
        ws = workbook["Cálculo Fiscal"]
        _, mapping = find_header(ws, {"rota", "codigo", "status"})
        final_price_field = next((name for name in ("preco final", "total c tributos") if name in mapping), None)
        if final_price_field is None:
            raise ValueError("Preço final não encontrado em Cálculo Fiscal")
        for values in ws.iter_rows(min_row=4, values_only=True):
            row = row_dict(values, mapping)
            route, code = text(row.get("rota")), product_code(row.get("codigo"))
            if not route or not code:
                continue
            calculations[(route, code)] = {
                "base_price": money(row.get("preco s imp")),
                "final_price": money(row.get(final_price_field)),
                "total_taxes": money(row.get("total tributos") if "total tributos" in mapping else row.get("ipi icms st")),
                "calculation_status": text(row.get("status")) or "UNKNOWN",
                "tax_breakdown": {
                    name: value for name, value in {
                        "ipi": number(row.get("ipi")),
                        "icms_proprio": number(row.get("icms proprio") if "icms proprio" in mapping else row.get("icms proprio info")),
                        "icms_st": number(row.get("icms st")), "pis": number(row.get("pis n d")),
                        "cofins": number(row.get("cofins n d")), "fcp": number(row.get("fcp n d"))
                    }.items() if value is not None
                }
            }

    base_prices: dict[tuple[str, str], dict[str, Any]] = {}
    route_prices: dict[tuple[str, str], dict[str, Any]] = {}
    for route in ("PR-PR", "SP-SP", "PR-SC"):
        sheet = f"LISTA {route}"
        if sheet not in workbook.sheetnames:
            continue
        ws = workbook[sheet]
        _, mapping = find_header(ws, {"codigo ips", "preco s imp", "preco c impostos"})
        for values in ws.iter_rows(min_row=3, values_only=True):
            row = row_dict(values, mapping)
            code = product_code(row.get("codigo ips"))
            base = money(row.get("preco s imp"))
            final = money(row.get("preco c impostos"))
            if not code:
                continue
            if base is not None:
                base_prices[(route[:2], code)] = {"base_price": base, "currency": "BRL"}
            calc = calculations.get((route, code), {})
            final = calc.get("final_price") if calc.get("final_price") is not None else final
            if final is None:
                continue
            fields = {"final_price": final, "currency": "BRL"}
            for name in ("base_price", "total_taxes", "tax_breakdown", "calculation_status"):
                put(fields, name, calc.get(name))
            route_prices[(route, code)] = fields
    for (branch, code), fields in base_prices.items():
        records.append({"area": "BASE_PRICE", "product_code": code, "branch_code": branch,
                        "fields": fields, "field_mask": list(fields)})
    for (route, code), fields in route_prices.items():
        records.append({"area": "ROUTE_PRICE", "product_code": code, "route": route,
                        "fields": fields, "field_mask": list(fields)})


def read_fiscal_bases(workbook) -> dict[str, list[dict[str, Any]]]:
    ncm_rules: list[dict[str, Any]] = []
    group_rules: list[dict[str, Any]] = []

    if "Dados Fiscais" in workbook.sheetnames:
        ws = workbook["Dados Fiscais"]
        _, mapping = find_header(ws, {"ncm", "uf origem", "uf destino", "mva sap", "icms inter", "icms interna", "ipi"})
        header_row = next(i for i, values in enumerate(ws.iter_rows(values_only=True), 1)
                          if {"ncm", "uf origem", "uf destino", "mva sap", "icms inter", "icms interna", "ipi"}
                          .issubset({key(value) for value in values}))
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            row = row_dict(values, mapping)
            ncm = digits(row.get("ncm"))
            origin = (text(row.get("uf origem")) or "").upper()
            destination = (text(row.get("uf destino")) or "").upper()
            if not ncm or len(ncm) != 8 or not re.fullmatch(r"[A-Z]{2}", origin) or not re.fullmatch(r"[A-Z]{2}", destination):
                continue
            rule: dict[str, Any] = {
                "rule_key": f"{ncm}|{origin}|{destination}",
                "ncm": ncm,
                "origin_state": origin,
                "destination_state": destination,
            }
            for source, target in (
                ("mva sap", "mva_rate"), ("icms inter", "interstate_icms_rate"),
                ("icms interna", "internal_icms_rate"), ("reducao bc", "base_reduction_rate"),
                ("ipi", "ipi_rate"), ("pis", "pis_rate"), ("cofins", "cofins_rate"),
                ("fcp", "fcp_rate"), ("frete", "freight_rate"),
                ("seguro", "insurance_rate"), ("outras desp", "other_expenses_rate"),
            ):
                put(rule, target, workbook_fraction_rate(row.get(source)))
            # The workbook's NCM fallback executes the ICMS-ST formula whenever
            # MVA is present (including an explicit zero). Group rules carry an
            # explicit SIM/NAO decision and take priority during calculation.
            rule["has_st"] = "mva_rate" in rule
            for source, target in (("cest", "cest"), ("cfop", "cfop"), ("csosn", "cst_code"),
                                   ("observacoes", "notes")):
                value = digits(row.get(source)) if source == "cest" else text(row.get(source))
                put(rule, target, value)
            ncm_rules.append(rule)

    if "Regras por Grupo" in workbook.sheetnames:
        ws = workbook["Regras por Grupo"]
        _, mapping = find_header(ws, {"ncm", "grupo de item", "rota", "mva derivada", "ipi correto",
                                      "icms inter", "icms interna", "st aplicavel"})
        header_row = next(i for i, values in enumerate(ws.iter_rows(values_only=True), 1)
                          if {"ncm", "grupo de item", "rota", "mva derivada", "ipi correto",
                              "icms inter", "icms interna", "st aplicavel"}
                          .issubset({key(value) for value in values}))
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            row = row_dict(values, mapping)
            ncm = digits(row.get("ncm"))
            item_group = text(row.get("grupo de item"))
            route = (text(row.get("rota")) or "").upper()
            if not ncm or len(ncm) != 8 or not item_group or not re.fullmatch(r"[A-Z]{2}-[A-Z]{2}", route):
                continue
            rule = {
                "rule_key": f"{ncm}|{item_group}|{route}",
                "ncm": ncm,
                "item_group": item_group,
                "route": route,
                "mva_rate": workbook_fraction_rate(row.get("mva derivada")),
                "ipi_rate": workbook_fraction_rate(row.get("ipi correto")),
                "interstate_icms_rate": workbook_fraction_rate(row.get("icms inter")),
                "internal_icms_rate": workbook_fraction_rate(row.get("icms interna")),
                "has_st": yes_no(row.get("st aplicavel")),
            }
            put(rule, "sample_base_price", money(row.get("base amostra")))
            put(rule, "sample_final_price", money(row.get("preco sap")))
            put(rule, "sample_st_amount", money(row.get("st alvo")))
            put(rule, "sample_product_code", product_code(row.get("codigo amostra")))
            group_rules.append(rule)

    if len({rule["rule_key"] for rule in ncm_rules}) != len(ncm_rules):
        raise ValueError("REGRAS_FISCAIS_NCM_DUPLICADAS")
    if len({rule["rule_key"] for rule in group_rules}) != len(group_rules):
        raise ValueError("REGRAS_FISCAIS_GRUPO_DUPLICADAS")
    return {"ncm_rules": ncm_rules, "group_rules": group_rules}


def build(source: Path, source_updated_at_override: str | None = None) -> dict[str, Any]:
    initial_stat = source.stat()
    digest = sha256_file(source)
    source_updated_at = source_updated_at_override or datetime.fromtimestamp(initial_stat.st_mtime, timezone.utc).isoformat()
    try:
        datetime.fromisoformat(source_updated_at.replace("Z", "+00:00"))
    except (TypeError, ValueError) as error:
        raise ValueError("DATA_ORIGEM_INVALIDA") from error
    workbook = load_workbook(source, read_only=True, data_only=True)
    products: dict[str, dict[str, Any]] = {}
    records: list[dict[str, Any]] = []
    fiscal_bases: dict[str, list[dict[str, Any]]] = {"ncm_rules": [], "group_rules": []}
    try:
        read_products(workbook, products)
        for record in products.values():
            record["field_mask"] = list(record["fields"])
            records.append(record)
        read_stock(workbook, records)
        read_prices(workbook, records)
        fiscal_bases = read_fiscal_bases(workbook)
    finally:
        workbook.close()
    # OneDrive may replace the workbook while it is being read. Never publish a
    # payload whose records came from one version but hash/mtime from another.
    final_digest = sha256_file(source)
    final_stat = source.stat()
    if (final_digest != digest or final_stat.st_mtime_ns != initial_stat.st_mtime_ns
            or final_stat.st_size != initial_stat.st_size):
        raise RuntimeError("EXCEL_ALTERADO_DURANTE_LEITURA")
    for index, record in enumerate(records, 1):
        record["row_number"] = index
        record["source_updated_at"] = source_updated_at
        record["source_version"] = digest
    return {
        "source_name": "Excel Mestre",
        "source_version": digest,
        "source_updated_at": source_updated_at,
        "file_hash": digest,
        "original_filename": source.name,
        "file_size": initial_stat.st_size,
        "records": records,
        "fiscal_bases": fiscal_bases,
        "summary": {
            "records": len(records),
            "products": sum(r["area"] == "PRODUCT" for r in records),
            "stocks": sum(r["area"] == "STOCK" for r in records),
            "base_prices": sum(r["area"] == "BASE_PRICE" for r in records),
            "route_prices": sum(r["area"] == "ROUTE_PRICE" for r in records),
            "fiscal_ncm_rules": len(fiscal_bases["ncm_rules"]),
            "fiscal_group_rules": len(fiscal_bases["group_rules"])
        }
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    payload = build(args.workbook.resolve())
    encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if args.output:
        args.output.write_text(encoded, encoding="utf-8")
    else:
        print(encoded)
    print(json.dumps(payload["summary"], ensure_ascii=False), file=__import__("sys").stderr)


if __name__ == "__main__":
    main()
